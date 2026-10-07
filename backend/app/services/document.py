"""Document ingestion and processing service."""
from __future__ import annotations
from app.rag.retrievers.retriever_manager import RetrieverManager

import hashlib
import codecs
import logging
import uuid
import zipfile
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path

import aiofiles
import anyio
import pymupdf
from docx import Document as DocxDocument
from pypdf import PdfReader
from unstructured.cleaners.core import clean
from unstructured.partition.text import partition_text
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.settings import get_settings
from app.core.ingestion_limits import (
    MAX_DOCUMENT_BYTES, MAX_DOCUMENT_CHUNKS, MAX_EXTRACTED_CHARACTERS,
    MAX_DOCX_EXPANDED_BYTES, MAX_DOCX_ENTRIES, MAX_PDF_PAGES,
    UPLOAD_READ_BYTES, document_size_limit,
)
from app.db.session import get_session_factory
from app.models.entities import Document, Embedding, SourceStatus
from app.schemas.document import (
    ChunkingConfig,
    DocumentChunk,
    DocumentFormat,
    DocumentMetadata,
    DocumentProcessingResult,
    EmbeddingConfig,
    ProcessingConfig,
)
from app.services.embedding_helper import generate_embeddings

# Supported MIME types and their formats
SUPPORTED_MIME_TYPES = {
    "application/pdf": DocumentFormat.pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocumentFormat.docx,
    "text/plain": DocumentFormat.txt,
    "text/csv": DocumentFormat.csv,
    "text/markdown": DocumentFormat.markdown,
    "text/x-markdown": DocumentFormat.markdown,
}

# File extensions to format mapping
EXTENSION_TO_FORMAT = {
    ".pdf": DocumentFormat.pdf,
    ".docx": DocumentFormat.docx,
    ".txt": DocumentFormat.txt,
    ".csv": DocumentFormat.csv,
    ".md": DocumentFormat.markdown,
    ".markdown": DocumentFormat.markdown,
}

MAX_FILE_SIZE = MAX_DOCUMENT_BYTES
INGEST_LIMITER = anyio.CapacityLimiter(2)
logger = logging.getLogger(__name__)


class DocumentValidationError(Exception):
    """Raised when document validation fails."""

    def __init__(self, message: str, code: str = "VALIDATION_ERROR"):
        self.message = message
        self.code = code
        super().__init__(message)


class DocumentProcessingError(Exception):
    """Raised when document processing fails."""

    def __init__(self, message: str, code: str = "PROCESSING_ERROR"):
        self.message = message
        self.code = code
        super().__init__(message)


def resolve_storage_key(root: Path, key: str) -> Path:
    """Resolve a relative storage key and reject escapes, including symlinks."""
    try:
        relative = Path(key)
        resolved_root = root.resolve()
        if not key or relative.is_absolute() or ".." in relative.parts or "\\" in key:
            raise ValueError("Invalid storage key")
        path = (resolved_root / relative).resolve()
        if path == resolved_root or not path.is_relative_to(resolved_root):
            raise ValueError("Storage key escapes root")
        return path
    except (OSError, ValueError, RuntimeError) as exc:
        raise DocumentValidationError(
            "Invalid document storage location", code="INVALID_STORAGE_KEY"
        ) from exc


def detect_format(file_path: Path, content_type: str | None = None) -> DocumentFormat:
    """
    Detect document format from file extension and MIME type.

    Args:
        file_path: Path to the file
        content_type: Optional MIME type from upload

    Returns:
        DocumentFormat enum value

    Raises:
        DocumentValidationError: If format is not supported
    """
    # First try MIME type from upload
    if content_type and content_type in SUPPORTED_MIME_TYPES:
        return SUPPORTED_MIME_TYPES[content_type]

    # Try file extension
    suffix = file_path.suffix.lower()
    if suffix in EXTENSION_TO_FORMAT:
        return EXTENSION_TO_FORMAT[suffix]

    # No recognized extension or MIME type - reject
    raise DocumentValidationError(
        "Unsupported file format. Supported formats: PDF, DOCX, TXT, CSV, Markdown",
        code="UNSUPPORTED_FORMAT",
    )


def validate_file(
    file_path: Path, content_type: str | None, size_bytes: int | None
) -> DocumentFormat:
    """
    Validate uploaded file.

    Args:
        file_path: Path to the uploaded file
        content_type: MIME type from upload
        size_bytes: File size in bytes

    Returns:
        DocumentFormat if valid

    Raises:
        DocumentValidationError: If validation fails
    """
    # Check file exists
    if not file_path.exists():
        raise DocumentValidationError("File not found", code="FILE_NOT_FOUND")

    # Check file size
    actual_size = file_path.stat().st_size
    if actual_size == 0:
        raise DocumentValidationError("File is empty", code="EMPTY_FILE")

    if actual_size > MAX_FILE_SIZE:
        raise DocumentValidationError(
            f"File size {actual_size} bytes exceeds maximum of {MAX_FILE_SIZE} bytes",
            code="FILE_TOO_LARGE",
        )

    if size_bytes is not None and size_bytes != actual_size:
        raise DocumentValidationError(
            f"Size mismatch: expected {size_bytes}, got {actual_size}",
            code="SIZE_MISMATCH",
        )

    # Detect and validate format
    format_ = detect_format(file_path, content_type)

    # Additional validation per format
    _validate_format_specific(file_path, format_)

    return format_


def _validate_format_specific(file_path: Path, format_: DocumentFormat) -> None:
    """Perform format-specific validation."""
    try:
        _validate_signature(file_path, format_)
        if format_ == DocumentFormat.pdf:
            # Try to open with both pymupdf and pypdf
            with pymupdf.open(str(file_path)) as doc:
                if doc.page_count == 0:
                    raise DocumentValidationError("PDF has no pages", code="EMPTY_PDF")
                if doc.page_count > MAX_PDF_PAGES:
                    raise DocumentValidationError("PDF has too many pages", code="EXTRACTION_LIMIT")

            # Also verify with pypdf
            reader = PdfReader(str(file_path))
            if len(reader.pages) == 0:
                raise DocumentValidationError(
                    "PDF has no pages (pypdf)", code="EMPTY_PDF"
                )

        elif format_ == DocumentFormat.docx:
            doc = DocxDocument(str(file_path))
            # Check if document has any content
            has_content = any(p.text.strip() for p in doc.paragraphs)
            if not has_content:
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                has_content = True
                                break
            if not has_content:
                raise DocumentValidationError("DOCX has no content", code="EMPTY_DOCX")

        elif format_ == DocumentFormat.csv:
            # Try to read as CSV
            import csv

            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                sniffer = csv.Sniffer()
                try:
                    dialect = sniffer.sniff(f.read(1024))
                    f.seek(0)
                    reader = csv.reader(f, dialect)
                    if next(reader, None) is None:
                        raise DocumentValidationError(
                            "CSV has no rows", code="EMPTY_CSV"
                        )
                except csv.Error:
                    raise DocumentValidationError(
                        "Invalid CSV format", code="INVALID_CSV"
                    )

    except DocumentValidationError:
        raise
    except Exception as e:
        raise DocumentValidationError(
            f"Format validation failed: {e}", code="FORMAT_VALIDATION_FAILED"
        )


def _validate_signature(file_path: Path, format_: DocumentFormat) -> None:
    """Inspect bytes before invoking a parser; MIME headers are only a hint."""
    with file_path.open("rb") as source:
        header = source.read(16)
        source.seek(0)
        if format_ == DocumentFormat.pdf:
            if not header.startswith(b"%PDF-"):
                raise DocumentValidationError("Invalid PDF signature", code="INVALID_SIGNATURE")
        elif format_ == DocumentFormat.docx:
            if not header.startswith(b"PK\x03\x04"):
                raise DocumentValidationError("Invalid DOCX signature", code="INVALID_SIGNATURE")
            with zipfile.ZipFile(source) as archive:
                entries = archive.infolist()
                if (len(entries) > MAX_DOCX_ENTRIES
                        or sum(entry.file_size for entry in entries) > MAX_DOCX_EXPANDED_BYTES):
                    raise DocumentValidationError("DOCX expansion exceeds limit", code="EXTRACTION_LIMIT")
                if not {"[Content_Types].xml", "word/document.xml"}.issubset(archive.namelist()):
                    raise DocumentValidationError("Invalid DOCX structure", code="INVALID_SIGNATURE")
        else:
            if header.startswith((b"\x7fELF", b"MZ", b"PK\x03\x04", b"%PDF-", b"\x1f\x8b")):
                raise DocumentValidationError("Binary file is not text", code="INVALID_SIGNATURE")
            decoder = codecs.getincrementaldecoder("utf-8")("strict")
            try:
                while data := source.read(UPLOAD_READ_BYTES):
                    text = decoder.decode(data)
                    if any((ord(char) < 32 and char not in "\t\r\n\f") or 127 <= ord(char) <= 159 for char in text):
                        raise DocumentValidationError("Binary file is not text", code="INVALID_SIGNATURE")
                decoder.decode(b"", final=True)
            except UnicodeDecodeError as exc:
                raise DocumentValidationError("Text documents must use UTF-8", code="INVALID_ENCODING") from exc


class _BoundedTextParts(list):
    """Stop accumulation of expanded PDF/DOCX/CSV text at a fixed ceiling."""

    def __init__(self):
        super().__init__()
        self.characters = 0

    def append(self, text):
        self.characters += len(text) + 2
        if self.characters > MAX_EXTRACTED_CHARACTERS:
            raise DocumentValidationError("Extracted text exceeds limit", code="EXTRACTION_LIMIT")
        super().append(text)


def compute_checksum(file_path: Path) -> str:
    """Compute SHA256 checksum of file."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha256.update(chunk)
    return f"sha256:{sha256.hexdigest()}"


def extract_text_pdf(file_path: Path) -> tuple[str, DocumentMetadata]:
    """Extract text and metadata from PDF using pymupdf."""
    doc = pymupdf.open(str(file_path))
    text_parts = _BoundedTextParts()
    metadata = DocumentMetadata(format=DocumentFormat.pdf)

    # Extract document metadata
    pdf_meta = doc.metadata
    if pdf_meta:
        metadata.title = pdf_meta.get("title")
        metadata.author = pdf_meta.get("author")
        metadata.subject = pdf_meta.get("subject")
        metadata.creator = pdf_meta.get("creator")
        metadata.producer = pdf_meta.get("producer")
        if pdf_meta.get("creationDate"):
            try:
                # Parse PDF date format: D:YYYYMMDDHHmmSSOHH'mm'
                date_str = pdf_meta["creationDate"]
                if date_str.startswith("D:"):
                    date_str = date_str[2:16]
                    metadata.creation_date = datetime.strptime(date_str, "%Y%m%d%H%M%S")
            except Exception:
                pass
        if pdf_meta.get("modDate"):
            try:
                date_str = pdf_meta["modDate"]
                if date_str.startswith("D:"):
                    date_str = date_str[2:16]
                    metadata.modification_date = datetime.strptime(
                        date_str, "%Y%m%d%H%M%S"
                    )
            except Exception:
                pass

    metadata.page_count = doc.page_count

    # Extract text from each page
    try:
        if doc.page_count > MAX_PDF_PAGES:
            raise DocumentValidationError("PDF has too many pages", code="EXTRACTION_LIMIT")
        for page_num in range(doc.page_count):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text)
    finally:
        doc.close()

    full_text = "\n\n".join(text_parts)
    metadata.word_count = len(full_text.split())
    metadata.character_count = len(full_text)

    return full_text, metadata


def extract_text_docx(file_path: Path) -> tuple[str, DocumentMetadata]:
    """Extract text and metadata from DOCX."""
    doc = DocxDocument(str(file_path))
    text_parts = _BoundedTextParts()
    metadata = DocumentMetadata(format=DocumentFormat.docx)

    # Extract core properties
    core_props = doc.core_properties
    if core_props:
        metadata.title = core_props.title
        metadata.author = core_props.author
        metadata.subject = core_props.subject
        metadata.creator = core_props.author
        if core_props.created:
            metadata.creation_date = core_props.created
        if core_props.modified:
            metadata.modification_date = core_props.modified

    # Extract text from paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            if row_text.strip():
                text_parts.append(row_text)

    full_text = "\n\n".join(text_parts)
    metadata.word_count = len(full_text.split())
    metadata.character_count = len(full_text)

    return full_text, metadata


def extract_text_txt(file_path: Path) -> tuple[str, DocumentMetadata]:
    """Extract text from plain text file."""
    metadata = DocumentMetadata(format=DocumentFormat.txt)

    # Try different encodings
    encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]
    text = None

    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        # Last resort: read as binary and decode with replacement
        with open(file_path, "rb") as f:
            text = f.read().decode("utf-8", errors="replace")

    metadata.word_count = len(text.split())
    metadata.character_count = len(text)

    return text, metadata


def extract_text_csv(file_path: Path) -> tuple[str, DocumentMetadata]:
    """Extract text from CSV file."""
    import csv

    metadata = DocumentMetadata(format=DocumentFormat.csv)
    text_parts = _BoundedTextParts()

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        # Detect dialect
        sample = f.read(1024)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(f, dialect)
        headers = next(reader, None)
        if headers:
            text_parts.append(" | ".join(headers))

        row_count = 0
        for row in reader:
            row_count += 1
            text_parts.append(" | ".join(row))

    metadata.custom_properties = {"row_count": row_count, "headers": headers}
    full_text = "\n".join(text_parts)
    metadata.word_count = len(full_text.split())
    metadata.character_count = len(full_text)

    return full_text, metadata


def extract_text_markdown(file_path: Path) -> tuple[str, DocumentMetadata]:
    """Extract text from Markdown file."""
    metadata = DocumentMetadata(format=DocumentFormat.markdown)

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    metadata.word_count = len(text.split())
    metadata.character_count = len(text)

    return text, metadata


def extract_text(
    file_path: Path, format_: DocumentFormat
) -> tuple[str, DocumentMetadata]:
    """Extract text and metadata from document based on format."""
    extractors = {
        DocumentFormat.pdf: extract_text_pdf,
        DocumentFormat.docx: extract_text_docx,
        DocumentFormat.txt: extract_text_txt,
        DocumentFormat.csv: extract_text_csv,
        DocumentFormat.markdown: extract_text_markdown,
    }

    extractor = extractors.get(format_)
    if not extractor:
        raise DocumentProcessingError(
            f"No extractor for format: {format_}", code="NO_EXTRACTOR"
        )

    return extractor(file_path)


def clean_text(text: str) -> str:
    """Clean extracted text using unstructured cleaners."""
    # Use unstructured's clean function
    cleaned = clean(
        text,
        bullets=True,
        extra_whitespace=True,
        dashes=True,
        trailing_punctuation=True,
    )
    return cleaned


def chunk_text(
    text: str,
    config: ChunkingConfig,
    metadata: DocumentMetadata | None = None,
) -> list[DocumentChunk]:
    """
    Chunk text into smaller pieces.

    Uses unstructured's chunk_by_title for semantic chunking,
    with fallback to fixed-size chunking.
    """
    chunks = []
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentValidationError("Extracted text exceeds limit", code="EXTRACTION_LIMIT")

    if config.strategy == "semantic":
        # Use unstructured's semantic chunking
        from unstructured.chunking.title import chunk_by_title

        elements = partition_text(text=text)
        chunked_elements = chunk_by_title(
            elements,
            max_characters=config.chunk_size * 4,  # Approximate chars per token
            overlap=config.chunk_overlap * 4,
        )

        for i, element in enumerate(chunked_elements):
            _check_chunk_limit(i + 1)
            chunk_text = str(element)
            chunks.append(
                DocumentChunk(
                    index=i,
                    content=chunk_text,
                    token_count=int(len(chunk_text.split()) * 1.3),  # Rough estimate
                    metadata={"strategy": "semantic"},
                )
            )

    elif config.strategy == "recursive":
        # Recursive character text splitter (similar to LangChain)
        from unstructured.chunking.basic import chunk_elements

        elements = partition_text(text=text)
        chunked_elements = chunk_elements(
            elements,
            max_characters=config.chunk_size * 4,
            overlap=config.chunk_overlap * 4,
        )

        for i, element in enumerate(chunked_elements):
            _check_chunk_limit(i + 1)
            chunk_text = str(element)
            chunks.append(
                DocumentChunk(
                    index=i,
                    content=chunk_text,
                    token_count=int(len(chunk_text.split()) * 1.3),
                    metadata={"strategy": "recursive"},
                )
            )

    else:  # fixed
        # Simple fixed-size chunking
        words = text.split()
        chunk_size_words = int(config.chunk_size * 0.75)  # Approximate words per token
        overlap_words = int(config.chunk_overlap * 0.75)

        for i in range(0, len(words), chunk_size_words - overlap_words):
            _check_chunk_limit(len(chunks) + 1)
            chunk_words = words[i : i + chunk_size_words]
            chunk_text = " ".join(chunk_words)
            chunks.append(
                DocumentChunk(
                    index=len(chunks),
                    content=chunk_text,
                    token_count=len(chunk_words),
                    metadata={"strategy": "fixed"},
                )
            )

    return chunks


def _check_chunk_limit(count: int) -> None:
    if count > MAX_DOCUMENT_CHUNKS:
        raise DocumentValidationError(
            f"Document exceeds the maximum of {MAX_DOCUMENT_CHUNKS} chunks",
            code="TOO_MANY_CHUNKS",
        )


@asynccontextmanager
async def _preserve_document_vectors(document: Document):
    """Compensate vector writes if indexing or SQL commit fails.

    Callers hold the document row lock until the SQL transaction completes.
    SQL and Chroma cannot share a transaction; process crashes still require
    reconciliation, but ordinary failures restore the prior vector snapshot.
    """
    document_id, user_id = str(document.id), str(document.user_id)

    def snapshot():
        retriever = RetrieverManager()
        return retriever, retriever.get_document_vectors(document_id, user_id)

    retriever, previous = await anyio.to_thread.run_sync(snapshot, limiter=INGEST_LIMITER)
    try:
        yield retriever
    except BaseException:
        def restore():
            retriever.delete_document_vectors(document_id, user_id)
            retriever.upsert_vectors(previous)

        try:
            with anyio.CancelScope(shield=True):
                await anyio.to_thread.run_sync(restore, limiter=INGEST_LIMITER)
        except Exception:
            logger.exception("Document vector recovery failed for %s", document_id)
        raise


async def delete_stored_document(session: AsyncSession, document: Document) -> None:
    """Remove vectors before committing SQL deletion; retain the file on failure."""
    root = Path(get_settings().upload_dir)
    resolve_storage_key(root, document.storage_key)
    storage_key = document.storage_key
    document_id, user_id = str(document.id), str(document.user_id)
    try:
        await session.delete(document)
        await session.flush()
        async with _preserve_document_vectors(document) as retriever:
            await anyio.to_thread.run_sync(
                retriever.delete_document_vectors, document_id, user_id, limiter=INGEST_LIMITER,
            )
            await session.commit()
    except Exception as exc:
        await session.rollback()
        raise DocumentProcessingError("Document deletion failed", code="DOCUMENT_DELETE_FAILED") from exc
    # SQL deletion is durable before physical cleanup. A cleanup failure must not
    # revive vectors or pretend that the committed SQL transaction rolled back.
    try:
        await anyio.to_thread.run_sync(lambda: resolve_storage_key(root, storage_key).unlink(missing_ok=True))
    except OSError:
        logger.exception("Deleted document file cleanup failed for %s", document_id)


async def store_document(
    document: Document,
    chunks: list[DocumentChunk],
    embeddings: list[list[float]],
    config: EmbeddingConfig,
) -> None:
    """Atomically replace a document's SQL chunks, serialized on its owner row."""
    _check_chunk_limit(len(chunks))
    if len(chunks) != len(embeddings) or len({c.index for c in chunks}) != len(chunks):
        raise DocumentValidationError(
            "Chunks must have unique indices and one embedding each",
            code="INVALID_EMBEDDINGS",
        )

    async with get_session_factory()() as session:

        db_document = await session.scalar(
            select(Document).where(
                Document.id == document.id, Document.user_id == document.user_id
            ).with_for_update()
        )

        if db_document is None:
            raise DocumentProcessingError(
                "Document not found",
                code="DOCUMENT_NOT_FOUND",
            )

        await session.execute(delete(Embedding).where(Embedding.document_id == document.id))
        session.add_all([
            Embedding(
                document_id=document.id,
                chunk_index=chunk.index,
                content=chunk.content,
                vector=embedding,
                model_name=config.model_name,
                token_count=chunk.token_count,
                metadata_=chunk.metadata,
            ) for chunk, embedding in zip(chunks, embeddings, strict=True)
        ])
        await session.flush()
        
        def index_chunks(retriever) -> None:
            retriever.delete_document_vectors(str(document.id), str(document.user_id))
            retriever.add_document_chunks({
                "documents": [chunk.content for chunk in chunks],
                "metadatas": [
                    {
                        "source_id": str(document.id),
                        "source_type": "document",
                        "document_id": str(document.id),
                        "user_id": str(document.user_id),
                        "filename": getattr(document, "filename", "unknown"),
                        "chunk_index": chunk.index,
                    }
                    for chunk in chunks
                ],
                "ids": [f"{document.id}_chunk_{chunk.index}" for chunk in chunks],
            })

        try:
            async with _preserve_document_vectors(db_document) as retriever:
                await anyio.to_thread.run_sync(index_chunks, retriever, limiter=INGEST_LIMITER)
                db_document.status = SourceStatus.ready
                await session.commit()
        except Exception as exc:
            await session.rollback()
            raise DocumentProcessingError(
                "Document vector indexing failed", code="RAG_INDEXING_FAILED"
            ) from exc


async def process_document(
    document_id: uuid.UUID,
    file_path: Path,
    format_: DocumentFormat,
    config: ProcessingConfig | None = None,
) -> DocumentProcessingResult:
    """
    Process a document through the full pipeline:
    validation -> cleaning -> chunking -> embedding -> storage
    """
    import time

    start_time = time.time()
    config = config or ProcessingConfig()

    async with get_session_factory()() as session:
        document = await session.get(Document, document_id)
        if not document:
            raise DocumentProcessingError(
                "Document not found", code="DOCUMENT_NOT_FOUND"
            )

        # Update status: pending (start processing)
        document.status = SourceStatus.pending
        await session.commit()

    try:
        # Always derive the input from the persisted key, never the caller's path.
        file_path = resolve_storage_key(Path(get_settings().upload_dir), document.storage_key)
        # Reprocessing must enforce the same limits as new uploads.
        format_ = await anyio.to_thread.run_sync(
            validate_file, file_path, document.mime_type, document.size_bytes,
            limiter=INGEST_LIMITER,
        )
        # Step 1: Extract text and metadata
        text, metadata = await anyio.to_thread.run_sync(
            extract_text, file_path, format_, limiter=INGEST_LIMITER
        )
        if len(text) > MAX_EXTRACTED_CHARACTERS:
            raise DocumentValidationError("Extracted text exceeds limit", code="EXTRACTION_LIMIT")

        # Step 2: Clean text
        if config.clean_text:
            text = await anyio.to_thread.run_sync(
                clean_text, text, limiter=INGEST_LIMITER
            )

        # Step 3: Chunk text
        chunks = await anyio.to_thread.run_sync(
            chunk_text, text, config.chunking, metadata, limiter=INGEST_LIMITER
        )

        # Step 4: Generate embeddings
        _check_chunk_limit(len(chunks))
        embeddings = await generate_embeddings(
            chunks, config.embedding.model_name, config.embedding.dimensions
        )

        # Step 5: Store in database
        await store_document(document, chunks, embeddings, config.embedding)

        processing_time_ms = int((time.time() - start_time) * 1000)

        return DocumentProcessingResult(
            document_id=document_id,
            filename=document.filename,
            format=format_,
            total_chunks=len(chunks),
            total_tokens=sum(c.token_count or 0 for c in chunks),
            processing_time_ms=processing_time_ms,
            chunks=chunks,
            metadata=metadata.model_dump() if config.extract_metadata else {},
        )

    except (DocumentProcessingError, DocumentValidationError):
        async with get_session_factory()() as session:
            failed_document = await session.get(Document, document_id)
            if failed_document:
                failed_document.status = SourceStatus.failed
                failed_document.metadata_ = {"error": "Document processing failed"}
                await session.commit()
        raise
    except Exception as e:
        # Update status to failed
        async with get_session_factory()() as session:
            document = await session.get(Document, document_id)
            if document:
                document.status = SourceStatus.failed
                document.metadata_ = {"error": "Document processing failed"}
                await session.commit()
        raise DocumentProcessingError(
            "Document processing failed", code="PROCESSING_FAILED"
        ) from e


async def upload_document(
    file_content: bytes,
    filename: str,
    content_type: str,
    user_id: uuid.UUID,
    config: ProcessingConfig | None = None,
) -> Document:
    """
    Upload and initiate processing of a document.

    Returns the created Document entity.
    """
    settings = get_settings()
    if len(file_content) > document_size_limit(getattr(settings, "max_file_size_mb", 15)):
        raise DocumentValidationError("File exceeds maximum size", code="FILE_TOO_LARGE")

    # Generate unique storage key
    document_id = uuid.uuid4()
    storage_key = f"uploads/{user_id}/{document_id}/content"
    root = Path(settings.upload_dir)
    file_path = resolve_storage_key(root, storage_key)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    created = False
    committed = False
    try:
        # Exclusive creation also prevents accidental replacement on key collision.
        async with aiofiles.open(resolve_storage_key(root, storage_key), "xb") as f:
            created = True
            await f.write(file_content)

        stored_path = resolve_storage_key(root, storage_key)
        await anyio.to_thread.run_sync(
            validate_file,
            stored_path,
            content_type,
            len(file_content),
            limiter=INGEST_LIMITER,
        )
        checksum = await anyio.to_thread.run_sync(
            compute_checksum, stored_path, limiter=INGEST_LIMITER
        )

        async with get_session_factory()() as session:
            document = Document(
                id=document_id,
                user_id=user_id,
                filename=filename,
                storage_key=storage_key,
                mime_type=content_type,
                size_bytes=len(file_content),
                checksum=checksum,
                status=SourceStatus.pending,
                metadata_={},
            )
            session.add(document)
            await session.commit()
            committed = True
            await session.refresh(document)
    finally:
        if created and not committed:
            resolve_storage_key(root, storage_key).unlink(missing_ok=True)

    return document
