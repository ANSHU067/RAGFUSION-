"""Regression tests for constrained Render workers and deployment configuration."""

import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from requests import Request
from requests.exceptions import ReadTimeout, ConnectionError

from app.config.settings import Settings
from app.db import session as database
from app.services import youtube_service as youtube
from app.services.document import chunk_text
from app.schemas.document import ChunkingConfig


def test_real_document_processing_never_imports_nlp_or_downloads_models(tmp_path):
    script = r'''
import importlib.abc
import sys
import socket
from pathlib import Path

class RejectNLP(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'unstructured', 'spacy', 'en_core_web_sm', 'torch'}:
            raise AssertionError('Unexpected NLP import: ' + fullname)

sys.meta_path.insert(0, RejectNLP())
def reject_network(*args, **kwargs):
    raise AssertionError('Parsing must not download resources')
socket.socket.connect = reject_network

import pymupdf
from docx import Document
from app.services.document import extract_text, clean_text, chunk_text
from app.schemas.document import DocumentFormat, ChunkingConfig

root = Path(sys.argv[1])
pdf_path = root / 'sample.pdf'
with pymupdf.open() as pdf:
    page = pdf.new_page()
    page.insert_text((72, 72), 'Lightweight PDF extraction works.')
    pdf.save(pdf_path)
docx_path = root / 'sample.docx'
doc = Document()
doc.add_paragraph('Lightweight DOCX extraction works.')
doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table text is retained.'
doc.save(docx_path)
(root / 'sample.txt').write_text('Lightweight UTF-8 text: café.', encoding='utf-8')
(root / 'sample.md').write_text('# Lightweight Markdown\n\nKeep **formatting**.', encoding='utf-8')
for suffix, fmt in [('pdf', DocumentFormat.pdf), ('docx', DocumentFormat.docx),
                    ('txt', DocumentFormat.txt), ('md', DocumentFormat.markdown)]:
    text, metadata = extract_text(root / ('sample.' + suffix), fmt)
    assert 'Lightweight' in text
    if suffix == 'docx':
        assert 'Table text is retained.' in text
    for strategy in ('recursive', 'semantic', 'fixed'):
        chunks = chunk_text(clean_text(text), ChunkingConfig(strategy=strategy), metadata)
        assert chunks and 'Lightweight' in chunks[0].content
assert not {'unstructured', 'spacy', 'en_core_web_sm', 'torch'}.intersection(sys.modules)
from app.services.chunking_service import chunk_html_semantic, chunk_text_recursive
from app.schemas.website import WebsiteChunkingConfig
config = WebsiteChunkingConfig()
assert chunk_html_semantic('<h1>Title</h1><p>Public text</p>', config, 'https://example.com')
assert chunk_text_recursive('Public website text.', config, 'https://example.com')
'''
    result = subprocess.run(
        [sys.executable, "-c", script, str(tmp_path)],
        cwd=Path(__file__).resolve().parents[1],
        env={**os.environ, "RAG_VECTOR_STORE_PATH": str(tmp_path / "vectors")},
        capture_output=True, text=True, timeout=30,
    )
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("strategy", ["semantic", "recursive"])
def test_unbroken_text_chunks_cover_source_with_bounded_overlap(strategy):
    text = "".join(chr(0x4E00 + i) for i in range(1000))
    chunks = chunk_text(text, ChunkingConfig(strategy=strategy, chunk_size=100, chunk_overlap=20))
    assert all(len(chunk.content) <= 400 for chunk in chunks)
    rebuilt = chunks[0].content + "".join(chunk.content[80:] for chunk in chunks[1:])
    assert rebuilt == text


@pytest.mark.parametrize("name", ["RAGFUSION_GROQ_API_KEY", "DOCPRO_GROQ_API_KEY", "GROQ_API_KEY"])
@pytest.mark.parametrize("value", ['  "gsk_test"  ', "  'gsk_test'  ", ' "  gsk_test  " '])
def test_groq_key_aliases_and_cleanup(monkeypatch, name, value):
    for alias in ("RAGFUSION_GROQ_API_KEY", "DOCPRO_GROQ_API_KEY", "GROQ_API_KEY"):
        monkeypatch.delenv(alias, raising=False)
    monkeypatch.setenv(name, value)
    settings = Settings(_env_file=None)
    assert settings.groq_api_key == "gsk_test"
    assert "gsk_test" not in repr(settings)


def test_groq_canonical_key_has_precedence(monkeypatch):
    monkeypatch.setenv("RAGFUSION_GROQ_API_KEY", "canonical")
    monkeypatch.setenv("DOCPRO_GROQ_API_KEY", "legacy")
    monkeypatch.setenv("GROQ_API_KEY", "generic")
    assert Settings(_env_file=None).groq_api_key == "canonical"


def test_pipeline_uses_cleaned_settings_key(monkeypatch):
    from app.rag.pipelines import rag_pipeline
    import langchain_groq
    llm = Mock()
    monkeypatch.setattr(langchain_groq, "ChatGroq", llm)
    monkeypatch.setattr(rag_pipeline, "get_settings", lambda: Settings(_env_file=None, groq_api_key=' "gsk_configured" '))
    monkeypatch.setattr(rag_pipeline, "get_rag_config", lambda: SimpleNamespace(temperature=0, max_tokens=100))
    monkeypatch.setattr(rag_pipeline, "EmbeddingManager", Mock())
    monkeypatch.setattr(rag_pipeline, "RetrieverManager", Mock())
    monkeypatch.setattr(rag_pipeline.RAGPipeline, "_build_graph", Mock())
    rag_pipeline.RAGPipeline()
    assert llm.call_args.kwargs["api_key"] == "gsk_configured"
    rag_pipeline.RAGPipeline(api_key=" 'gsk_override' ")
    assert llm.call_args.kwargs["api_key"] == "gsk_override"


def test_postgres_connection_gets_cold_start_budget_without_changing_sqlite(monkeypatch):
    create = Mock()
    monkeypatch.setattr(database, "create_async_engine", create)
    monkeypatch.setattr(database, "get_settings", lambda: SimpleNamespace(database_connect_timeout_seconds=20))
    database.get_async_engine.cache_clear()
    try:
        database.get_async_engine("postgresql://test:test@localhost/example")
        args = create.call_args.kwargs
        assert args["pool_pre_ping"] is True
        assert args["connect_args"]["timeout"] == 20
        assert args["pool_timeout"] >= 10
        database.get_async_engine("sqlite+aiosqlite:///:memory:")
        assert create.call_args.kwargs["connect_args"] == {"check_same_thread": False}
    finally:
        database.get_async_engine.cache_clear()


def test_database_timeout_cannot_be_configured_below_ten_seconds():
    with pytest.raises(ValueError):
        Settings(_env_file=None, database_connect_timeout_seconds=1)


@pytest.mark.parametrize("failure", [ReadTimeout, ConnectionError])
def test_transcript_retries_network_errors_once_with_browser_headers(monkeypatch, failure):
    clock = [0.0]
    calls = []
    def send(self, request, **kwargs):
        calls.append(kwargs["timeout"])
        assert "Mozilla/5.0" in request.headers["User-Agent"]
        if len(calls) == 1:
            clock[0] += 20
            raise failure("Temporary upstream error")
        return "transcript response"
    monkeypatch.setattr(youtube.time, "monotonic", lambda: clock[0])
    monkeypatch.setattr(youtube.time, "sleep", lambda _: None)
    monkeypatch.setattr(youtube.HTTPSession, "send", send)
    with youtube.TranscriptSession(60) as session:
        request = session.prepare_request(Request("GET", "https://www.youtube.com/watch?v=test"))
        assert session.send(request) == "transcript response"
    assert len(calls) == 2
    assert [timeout.read_timeout for timeout in calls] == [20, 20]
    assert [timeout.total for timeout in calls] == [60, 40]


def test_transcript_retry_budget_is_shared_and_deadline_stops_requests(monkeypatch):
    send = Mock(side_effect=ReadTimeout("unavailable"))
    monkeypatch.setattr(youtube.HTTPSession, "send", send)
    monkeypatch.setattr(youtube.time, "sleep", lambda _: None)
    with youtube.TranscriptSession(youtube.time.monotonic() + 60) as session:
        with pytest.raises(ReadTimeout):
            session.send(object())
        assert send.call_count == 2
        with pytest.raises(ReadTimeout):
            session.send(object())
        assert send.call_count == 3
        session.deadline = 0
        with pytest.raises(youtube.HTTPTimeout):
            session.send(object())
        assert send.call_count == 3
